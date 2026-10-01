# -*- coding: utf-8 -*-
"""
国内批复核对工具 - Streamlit 版本

主要特性：
- 软换行（Shift+Enter）自动拆分成独立段落
- 外机两种格式识别（已批 / 未批）
- 用途映射：FERRY → N/M；BUSINESS → U/H
- 待申请简洁格式：`VPCSZ ZGSZ-ZBAD 07OCT 待申请`
- 追加文字（待取消/待变更/待申请）红字 + 黄底
- 三态判定：是 / 否 / 待确认
- 文本覆盖率预检查（< 60% 报错，60~85% 警告，≥ 85% 通过）
"""

import io
import re
import csv
import copy
import datetime
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

# =========================================================
# 常量
# =========================================================
MONTHS = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4,
    "MAY": 5, "JUN": 6, "JUL": 7, "AUG": 8,
    "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}

RED = "FF0000"
GREEN = "00B050"
HIGHLIGHT_YELLOW = "yellow"

MAX_CROSS_DAY_GAP_MIN = 600
EARLY_GREEN_THRESHOLD_MIN = 600

COVERAGE_ERROR_THRESHOLD = 0.60   # < 60% → 报错
COVERAGE_WARN_THRESHOLD = 0.85    # 60%~85% → 警告；≥ 85% → 通过

PILOT_RAW = """P001,庚凡,gengfan@amber-aviation.com
P002,张永一,zhangyongyi@amber-aviation.com
P003,梅峰,fmei@amber-aviation.com
P004,王斌,wangbin@amber-aviation.com
P019,"HEALY, Darran William",darranhealy@amber-aviation.com
P020,"BEEBE, Thaddeus John",thaddeusbeebe@amber-aviation.com
P032,林毅,ericlin@amber-aviation.com
P035,"Peter Robert, JACKSON",prjackson@amber-aviation.com
P036,王少雄,warrenwang@amber-aviation.com
P038,苗旺旺,johnmiao@amber-aviation.com
P039,"Yiftah, RAUCH",yiftahrauch@amber-aviation.com
P044,李辛欣,rockli@amber-aviation.com
P046,赵岩松,zhyszhao@amber-aviation.com
P051,彭罡,eugene.peng@humbleholding.com
P052,胡君量,brian.wu@humbleholding.com
P053,"Bruce Roderick, WAINES",brwaines@amber-aviation.com
P054,"Rodolfo, BONETTI",rbonetti@amber-aviation.com
P056,"Keith Robert, SHERREN",krsherren@amber-aviation.com
P057,"Oliver Viktor, RACZ",ovracz@amber-aviation.com
P059,蔡国俊,kctsai@amber-aviation.com
P061,李庆宏,qhli@amber-aviation.com
P065,宋炜,wsong@amber-aviation.com
P068,昝昭君,zjzan@amber-aviation.com
P069,"ROEDER, SIMONE ELKE",simoneroeder@amber-aviation.com
P070,"Herve Daniel, STAMM",hdstamm@amber-aviation.com
P071,孙浩,jasonsun@amber-aviation.com
P072,朱正宇,zyzhu@amber-aviation.com
P074,金尚明,smjin@amber-aviation.com
P075,"Eduard Pascal, Roski",eduardroski@amber-aviation.com
P077,刘凯,andyliu@amber-aviation.com
P078,张帆,fzhang@amber-aviation.com
P079,魏思远,wesleywei@amber-aviation.com
P080,刘爽,sliu@amber-aviation.com
P081,吴鹏,richardwu@amber-aviation.com
P082,刘汇川,frankliu@amber-aviation.com
P083,尤欣,xyou@amber-aviation.com
P084,李亚民,ymli@amber-aviation.com
P085,赵镭,lzhao@amber-aviation.com
P086,张贺新,hxzhang@amber-aviation.com
P087,孙赫,hesun@amber-aviation.com
P088,马坚,harryma@amber-aviation.com
P089,李晓龙,xlli@amber-aviation.com
P090,黄海东,hdhuang@amber-aviation.com
P091,马洪双,mikema@amber-aviation.com
PJZ001,张哲,zzhang@amber-aviation.com
PJZ002,郭春旭,charlesguo@amber-aviation.com
PJZ004,王国勤,leowang@amber-aviation.com
PJZ005,王莹,evawang@amber-aviation.com
PJZ007,徐卓,frankxu@amber-aviation.com
PJZ008,杨华,ariayang@amber-aviation.com
W070,王彦海,wang_yanhai@163.com
W213,沈志伟,cshum@tagaviation.com
W267,"Nathon Andrew G, NORBERG",naten7@hotmail.com
W268,"Daniel, RICHTER",pilotlocalizer@gmail.com
W270,杨涛,yang_tao2005@aliyun.com
W272,"Andrew Nigel, KING",Andrew.king@aero.bombardier.com
"""

AIRCRAFT_TYPE_MAP = {
    "B3926": "LJ60", "B652R": "GLF4", "B8105": "GLEX", "B8160": "GLF5",
    "B8262": "GLF4", "B8292": "GLF5", "B8309": "GLF5", "MLLIN": "GLEX",
    "N2QE": "GL5T", "N328LM": "GL7T", "N550DR": "GLF5", "N577QT": "F900",
    "N7777U": "GLEX", "N777ZH": "GLF5", "N88AY": "GLF5", "T7178HT": "GL7T",
    "T7CJK": "GLEX", "VPCSZ": "GL7T", "VPCVA": "GLF6", "B652Q": "GLF4",
    "B652S": "GLF4", "B65AP": "GLF4",
}

FERRY_KEYWORDS = ("调机", "维修")

st.set_page_config(page_title="国内批复核对工具", page_icon="✈️", layout="wide")


# =========================================================
# 工具
# =========================================================
def parse_date_token(token):
    token = token.strip().upper()
    day = int(token[:2])
    mon = MONTHS[token[2:5]]
    year_str = token[5:]
    year = 2000 + int(year_str) if len(year_str) == 2 else int(year_str)
    return datetime.date(year, mon, day)


def parse_hhmm(token):
    token = str(token).strip().zfill(4)
    return datetime.time(int(token[:2]), int(token[2:]))


def parse_hhmm_str(s):
    if not s:
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})$", str(s).strip())
    if not m:
        return None
    return datetime.time(int(m.group(1)), int(m.group(2)))


def is_b_reg(reg):
    return str(reg).strip().upper().startswith("B")


def to_beijing_datetime(reg, date_obj, hhmm_token):
    dt = datetime.datetime.combine(date_obj, parse_hhmm(hhmm_token))
    if not is_b_reg(reg):
        dt += datetime.timedelta(hours=8)
    return dt


def fmt_time(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    s = str(value).strip()
    m = re.match(r"(\d{1,2}):(\d{2})", s)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else s


def fmt_date(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def has_chinese(text):
    return bool(re.search(r"[\u4e00-\u9fff]", str(text)))


def time_diff_minutes(t1, t2):
    def to_min(t):
        h, m = t.split(":")
        return int(h) * 60 + int(m)
    try:
        return abs(to_min(t1) - to_min(t2))
    except Exception:
        return 9999


def hhmm_to_minutes(t):
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def flight_duration_minutes(dep_t, arr_t):
    try:
        d = hhmm_to_minutes(dep_t)
        a = hhmm_to_minutes(arr_t)
    except Exception:
        return None
    if a < d:
        a += 1440
    return a - d


def fmt_duration(mins):
    if mins is None:
        return ""
    sign = "-" if mins < 0 else ""
    m = abs(mins)
    return f"{sign}{m // 60}:{m % 60:02d}"


def compute_real_minute_diff(ap_date, ap_time_str, xl_date, xl_time_str):
    if ap_date is None or xl_date is None:
        return None
    try:
        ap_min = hhmm_to_minutes(ap_time_str)
        xl_min = hhmm_to_minutes(xl_time_str)
    except Exception:
        return None
    ap_abs = ap_date.toordinal() * 1440 + ap_min
    xl_abs = xl_date.toordinal() * 1440 + xl_min
    return ap_abs - xl_abs


@st.cache_data
def load_pilots():
    pilots = {}
    for line in PILOT_RAW.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            parts = next(csv.reader([line]))
        except Exception:
            parts = line.split(",")
        if len(parts) >= 2:
            pilots[parts[0].strip()] = parts[1].strip().strip('"')
    return pilots


def crew_all_chinese(crew_codes, pilots):
    pilot_codes = [c.strip() for c in crew_codes if c.strip().startswith(("P", "W"))]
    if not pilot_codes:
        return None
    for code in pilot_codes:
        if code not in pilots:
            return None
        if not has_chinese(pilots[code]):
            return False
    return True


def is_ferry_use(use_text):
    return any(k in use_text for k in FERRY_KEYWORDS)


# =========================================================
# 解析批复
# =========================================================
APPROVAL_RE = re.compile(
    r"^(?P<reg>[A-Z0-9\-]+)\s+"
    r"(?P<second>[A-Z0-9]+)\s+"
    r"(?P<dep>[A-Z]{4})(?P<dep_time>\d{4})\s+"
    r"(?P<arr_time>\d{4})(?P<arr>[A-Z]{4})\s+"
    r"ON\s+(?P<date>\d{2}[A-Z]{3}\d{2,4})\s+"
    r"(?P<rest>.+)$",
    re.IGNORECASE,
)


def parse_approval_line(text):
    m = APPROVAL_RE.match(text.strip())
    if not m:
        return None

    reg = m.group("reg").upper().replace("-", "")
    second = m.group("second").upper()
    dep = m.group("dep").upper()
    arr = m.group("arr").upper()
    dep_raw = m.group("dep_time")
    arr_raw = m.group("arr_time")
    date_raw = m.group("date").upper()
    rest = m.group("rest").strip()

    if is_b_reg(reg):
        ac_type = second
        flight_no = ""
    else:
        if second == reg:
            flight_no = second
            ac_type = ""
        else:
            ac_type = second
            flight_no = ""

    service, remark = "", ""
    sm = re.match(r"^(U/H|N/M)\s*(.*)$", rest, re.IGNORECASE)
    if sm:
        service = sm.group(1).upper()
        remark = sm.group(2).strip()
    else:
        upper = rest.upper()
        if upper.startswith("FERRY"):
            service = "N/M"
            remark = rest[5:].strip(" -–—\t")
        elif upper.startswith("BUSINESS"):
            service = "U/H"
            remark = rest[8:].strip(" -–—\t")
        else:
            parts = rest.split(None, 1)
            service = parts[0].upper() if parts else ""
            remark = parts[1].strip() if len(parts) > 1 else ""

    date_obj = parse_date_token(date_raw)
    return {
        "raw": text.strip(),
        "reg": reg,
        "type": ac_type,
        "flight_no": flight_no,
        "is_domestic": is_b_reg(reg),
        "dep": dep,
        "dep_time_raw": dep_raw,
        "arr_time_raw": arr_raw,
        "arr": arr,
        "date_raw": date_raw,
        "dep_dt_bj": to_beijing_datetime(reg, date_obj, dep_raw),
        "arr_dt_bj": to_beijing_datetime(reg, date_obj, arr_raw),
        "service": service,
        "remark": remark,
    }


def iter_doc_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for np in ncell.paragraphs:
                                yield np


# =========================================================
# 软换行（<w:br/>）拆分成独立段落
# =========================================================
def _collect_all_paragraphs(doc):
    result = []
    for p in doc.paragraphs:
        result.append(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    result.append(p)
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for np in ncell.paragraphs:
                                result.append(np)
    return result


def _split_paragraph_by_br(p_elem):
    parent = p_elem.getparent()
    if parent is None:
        return
    idx_in_parent = list(parent).index(p_elem)
    pPr = p_elem.find(qn('w:pPr'))

    groups = [[]]
    for child in list(p_elem):
        if child.tag == qn('w:pPr'):
            continue
        if child.tag == qn('w:r'):
            brs = child.findall(qn('w:br'))
            if not brs:
                groups[-1].append(copy.deepcopy(child))
            else:
                current_r_children = []
                for rc in list(child):
                    if rc.tag == qn('w:br'):
                        if current_r_children:
                            new_r = OxmlElement('w:r')
                            for x in current_r_children:
                                new_r.append(copy.deepcopy(x))
                            groups[-1].append(new_r)
                            current_r_children = []
                        groups.append([])
                    else:
                        current_r_children.append(rc)
                if current_r_children:
                    new_r = OxmlElement('w:r')
                    for x in current_r_children:
                        new_r.append(copy.deepcopy(x))
                    groups[-1].append(new_r)
        else:
            groups[-1].append(copy.deepcopy(child))

    if len(groups) <= 1:
        return

    parent.remove(p_elem)
    for i, group in enumerate(groups):
        new_p = OxmlElement('w:p')
        if pPr is not None:
            new_p.append(copy.deepcopy(pPr))
        for child in group:
            new_p.append(child)
        parent.insert(idx_in_parent + i, new_p)


def normalize_soft_breaks(doc):
    for p in _collect_all_paragraphs(doc):
        _split_paragraph_by_br(p._element)


# =========================================================
# Excel
# =========================================================
def load_excel_rows_from_bytes(data: bytes):
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["航段(北京时)"] if "航段(北京时)" in wb.sheetnames else wb.active

    rows = []
    for r in range(3, ws.max_row + 1):
        reg = ws.cell(r, 3).value
        if not reg:
            continue
        rows.append({
            "_idx": len(rows),
            "reg": str(reg).strip().upper(),
            "use": str(ws.cell(r, 4).value or "").strip(),
            "dep_date": fmt_date(ws.cell(r, 7).value),
            "dep_time": fmt_time(ws.cell(r, 8).value),
            "dep": str(ws.cell(r, 11).value or "").strip().upper(),
            "dep_city": str(ws.cell(r, 12).value or "").strip(),
            "arr": str(ws.cell(r, 13).value or "").strip().upper(),
            "arr_city": str(ws.cell(r, 14).value or "").strip(),
            "arr_date": fmt_date(ws.cell(r, 15).value),
            "arr_time": fmt_time(ws.cell(r, 16).value),
        })
    return rows


# =========================================================
# 文本航班信息
# =========================================================
FLIGHT_HEADER_RE = re.compile(
    r"^([A-Z0-9]+)\s+(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})(?:\s*\+1)?$"
)


def load_text_flights(text: str):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    flights = []
    i = 0
    pending_f = False

    while i < len(lines):
        line = lines[i]

        if line.upper() == "F":
            pending_f = True
            i += 1
            continue

        if line.upper() in ("TBA",):
            i += 1
            continue

        m = FLIGHT_HEADER_RE.match(line)
        if m:
            reg = m.group(1).upper()
            dep_time, arr_time = m.group(2), m.group(3)
            if i + 1 < len(lines):
                cm = re.match(r"^(.+?)\s+-\s+(.+)$", lines[i + 1])
                if cm:
                    crew = []
                    if i + 2 < len(lines):
                        cl = lines[i + 2].replace(" ", "")
                        if re.match(r"^[A-Z0-9,]+$", cl):
                            crew = [x for x in cl.split(",") if x]

                    flights.append({
                        "_idx": len(flights),
                        "reg": reg,
                        "dep_time": dep_time,
                        "arr_time": arr_time,
                        "dep_city": cm.group(1).strip(),
                        "arr_city": cm.group(2).strip(),
                        "crew": crew,
                        "is_ferry": pending_f,
                    })
                    pending_f = False
                    i += 3
                    continue
        i += 1
    return flights


# ===== 文本覆盖率检查（粘贴文本后立即执行，不需要点按钮）=====
if docx_file and excel_file and text_content:
    preview_excel_rows = load_excel_rows_from_bytes(excel_file.getvalue())
    preview_text_flights = load_text_flights(text_content)

    preview_city_to_icao = {}
    for row in preview_excel_rows:
        if row["dep_city"]:
            preview_city_to_icao[row["dep_city"]] = row["dep"]
        if row["arr_city"]:
            preview_city_to_icao[row["arr_city"]] = row["arr"]

    total, covered, missing = check_text_coverage(
        preview_excel_rows, preview_text_flights, preview_city_to_icao
    )

    st.subheader("🔍 文本覆盖率检查")

    if total == 0:
        st.warning("Excel 里没有国内航段（Z 开头机场），无需核对。")
    else:
        coverage = covered / total

        if coverage >= COVERAGE_WARN_THRESHOLD:
            st.success(
                f"✅ 文本计划覆盖率：{covered}/{total}（{coverage*100:.1f}%）"
            )
        elif coverage >= COVERAGE_ERROR_THRESHOLD:
            st.warning(
                f"⚠️ 文本计划覆盖率：{covered}/{total}（{coverage*100:.1f}%），"
                f"缺 {total - covered} 条。可以核对，但建议补全。"
            )
        else:
            st.error(
                f"❌ 文本计划覆盖率过低：{covered}/{total}（{coverage*100:.1f}%），"
                f"缺 {total - covered} 条。请先补全文本。"
            )

        # 只要有缺失就列出来，方便直接修改
        if missing:
            with st.expander(
                f"📋 缺失航段明细（{len(missing)} 条）—— 点开查看 / 复制",
                expanded=(coverage < COVERAGE_WARN_THRESHOLD),
            ):
                # 按日期 + 注册号排序，方便对照
                missing_sorted = sorted(
                    missing,
                    key=lambda r: (
                        r["dep_date"] or datetime.date.min,
                        r["reg"],
                        r["dep_time"],
                    ),
                )
                lines = []
                for r in missing_sorted:
                    date_str = r["dep_date"].strftime("%m-%d") if r["dep_date"] else "??-??"
                    lines.append(
                        f"{r['reg']}  {date_str}  {r['dep']}-{r['arr']}  "
                        f"{r['dep_time']}-{r['arr_time']}  ({r['dep_city']} → {r['arr_city']})"
                    )
                # 用代码块显示，方便整体复制
                st.code("\n".join(lines), language=None)

# ===== 正式核对（需要点击按钮）=====
if st.button("🚀 开始核对", type="primary", disabled=not (docx_file and excel_file and text_content)):
    with st.spinner("正在核对..."):
        try:
            rows, out_buf = run_check(
                docx_file.getvalue(),
                excel_file.getvalue(),
                text_content,
                load_pilots(),
            )
        except Exception as e:
            st.exception(e)
            st.stop()

    df = pd.DataFrame(rows)
    total_rows = len(df)
    diff_count = (df["是否一致"] == "否").sum() if total_rows else 0
    pending_count = (df["是否一致"] == "待确认").sum() if total_rows else 0

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("批复条数", total_rows)
    c2.metric("一致", total_rows - diff_count - pending_count)
    c3.metric("有差异", diff_count)
    c4.metric("待确认", pending_count)

    st.subheader("📋 核对结果")
    if total_rows == 0:
        st.warning("未在 docx 中识别到任何批复行。")
    else:
        def highlight(row):
            if row["是否一致"] == "待确认":
                return ["background-color: #fff3cd"] * len(row)
            if row["是否一致"] == "否":
                if "待变更" in str(row["备注"]):
                    return ["background-color: #e5f0ff"] * len(row)
                return ["background-color: #ffe5e5"] * len(row)
            return [""] * len(row)

        st.dataframe(
            df.style.apply(highlight, axis=1),
            use_container_width=True,
            hide_index=True,
        )

        st.subheader("🚨 差异明细")
        diffs_df = df[df["是否一致"] != "是"][["批复", "差异", "备注", "是否一致"]]
        if diffs_df.empty:
            st.success("✅ 所有批复与计划一致，未发现差异。")
        else:
            for _, r in diffs_df.iterrows():
                if r["是否一致"] == "待确认":
                    note_html = (
                        " <span style='color:#d97706;font-weight:bold'>"
                        f"【{r['备注']}】</span>"
                    )
                elif "待变更" in str(r["备注"]):
                    note_html = (
                        " <span style='color:#0066cc;font-weight:bold'>"
                        f"【{r['备注']}】</span>"
                    )
                elif r["备注"]:
                    note_html = (
                        f" <span style='color:red;font-weight:bold'>【{r['备注']}】</span>"
                    )
                else:
                    note_html = ""
                st.markdown(f"**`{r['批复']}`**{note_html}", unsafe_allow_html=True)
                if r["差异"] != "无":
                    for line in r["差异"].split("；"):
                        st.markdown(f"- {line}")

    st.subheader("📥 下载标红后的批复汇总表")
    st.download_button(
        label=f"下载 {docx_file.name}",
        data=out_buf.getvalue(),
        file_name=docx_file.name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
